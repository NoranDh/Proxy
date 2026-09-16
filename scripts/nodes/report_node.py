from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from langchain_openai import ChatOpenAI
from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from scripts.state import ProxyState


REPORT_SECTIONS = [
    ("1. Problem Description", [
        "Clear description of the issue (what failed)",
        "Equipment / tag number",
        "Failure type (leak, vibration, crack, overheating, etc.)",
        "Failure start time (date and operating hours)",
        "Current status (running, shutdown, degraded)",
    ]),
    ("2. Operating Conditions at Time of Failure", [
        "Pressure",
        "Temperature",
        "Flow rate",
        "Load / capacity (%)",
        "Process condition (normal / upset / transient)",
        "Trend data (last 24h / 7 days)",
        "Any abnormal spikes or trips",
    ]),
    ("3. Equipment Information", [
        "Equipment type (pump, exchanger, flare, etc.)",
        "Manufacturer and model",
        "Material of construction",
        "Design limits (pressure, temperature, corrosion allowance)",
        "P&ID reference (if available)",
        "Isometric drawings (if available)",
    ]),
    ("4. Failure History & Maintenance Records", [
        "Previous similar failures",
        "Last maintenance date",
        "Type of maintenance performed",
        "Replacement parts history",
        "Mean Time Between Failures (MTBF)",
    ]),
    ("5. Inspection & NDT Data", [
        "Ultrasonic Thickness (UT) readings",
        "Radiography Testing (RT)",
        "Magnetic Particle Testing (MT)",
        "Dye Penetrant Testing (PT)",
        "Eddy Current Testing",
        "Hardness testing",
        "Before/after comparison data (if available)",
    ]),
    ("6. Visual Evidence (Photos)", [
        "Clear images of failure location",
        "Close-up and wide-angle shots",
        "Thermal images (if available)",
    ]),
    ("7. Process & Chemical Data", [
        "Fluid composition",
        "pH levels",
        "Contaminants (H2S, CO2, chlorides)",
        "Water content",
    ]),
    ("8. Instrumentation & Control Data", [
        "Alarm history",
        "Trip logs",
        "Sensor readings",
        "Control valve positions",
    ]),
    ("9. Environmental Conditions", [
        "Ambient temperature",
        "Humidity",
        "Marine/saline exposure",
        "Dust/sand conditions",
    ]),
    ("10. Human & Operational Factors", [
        "Operator actions before failure",
        "SOP deviations",
        "Startup/shutdown issues",
    ]),
    ("11. Additional Notes", [
        "Temporary fixes applied",
        "Constraints (budget, shutdown limitations)",
        "Urgency level",
    ]),
    ("12. Root Cause Analysis", [
        "Confirmed mechanism",
        "Root cause",
        "RCA confidence",
        "5 Whys chain",
    ]),
    ("13. Recommendations & Action Plan", [
        "Immediate corrective actions",
        "Short-term recommendations",
        "Long-term prevention measures",
        "Action plan with owners and timelines",
    ]),
]


def report_node(state: ProxyState) -> dict:
    print("Report node is running...")

    llm = ChatOpenAI(model="gpt-4o", temperature=0)
    report_payload = _build_report_payload(state)

    prompt = f"""
You are a failure analysis engineer writing a root cause analysis report for an industrial incident at a Saudi Arabian oil and gas facility.

Fill every section and item from the provided template requirements.
Use the available evidence below. If a field is not supported by evidence, write "Not provided".
Be factual, concise, and engineering-focused.

Template requirements:
{json.dumps(REPORT_SECTIONS, indent=2, ensure_ascii=False)}

Evidence package:
{json.dumps(report_payload, indent=2, ensure_ascii=False)}

Return the final report as plain text using this exact structure:
- One section heading at a time, matching the template wording exactly (e.g. "1. Problem Description")
- Under each section, include each item as "Field name: value"
- Use the exact field names from the template
- For any missing information write "Field name: Not provided"
"""

    response = llm.invoke([{"role": "user", "content": prompt}])
    report_content = _clean_report_text(response.content)
    report_path = _save_report(state, report_content)

    return {
        "report_content": report_content,
        "report_path": str(report_path),
        "report_ready": True,
        "response_ready": True,
        "stage": "closed",
    }


def _build_report_payload(state: ProxyState) -> dict[str, Any]:
    rca_whys = state.get("rca_whys") or []
    whys_formatted = [
        f"Why {w['level']}: {w['question']} → {w.get('answer', 'Not answered')}"
        for w in rca_whys
    ]

    kb_recommendations = state.get("kb_recommendations") or []
    action_plan = state.get("action_plan") or []

    return {
        "case_id":                  state.get("case_id"),
        "equipment":                state.get("equipment"),
        "location":                 state.get("location"),
        "severity":                 state.get("severity"),
        "failure_type":             state.get("failure_type"),
        "raw_input":                state.get("raw_input"),
        "unified_context":          state.get("unified_context", ""),
        "dept_summaries":           state.get("dept_summaries") or {},
        "failure_mode_candidates":  state.get("failure_mode_candidates") or [],
        "confirmed_mechanism":      state.get("confirmed_mechanism", ""),
        "mechanism_reasoning":      state.get("mechanism_reasoning", ""),
        "root_cause":               state.get("root_cause"),
        "rca_confidence":           state.get("rca_confidence"),
        "rca_whys_chain":           whys_formatted,
        "review_decision":          state.get("review_decision"),
        "review_feedback":          state.get("review_feedback"),
        "kb_recommendations":       kb_recommendations,
        "action_plan":              action_plan,
        "mitigation_plan":          state.get("mitigation_plan"),
    }


def _clean_report_text(text: str) -> str:
    cleaned = text.strip()
    cleaned = re.sub(r"^```[a-zA-Z0-9_-]*\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    return cleaned.strip()


def _save_report(state: ProxyState, report_content: str) -> Path:
    root = Path(__file__).resolve().parents[2]
    reports_dir = root / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    case_id = state.get("case_id") or "unknown_case"
    safe_case_id = re.sub(r"[^A-Za-z0-9._-]+", "_", case_id)
    report_path = reports_dir / f"{safe_case_id}_rca_report.docx"

    doc = DocxDocument()

    # Title
    title = doc.add_heading("Root Cause Analysis Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata line
    meta = doc.add_paragraph(
        f"Case ID: {state.get('case_id', 'N/A')}  |  "
        f"Equipment: {state.get('equipment', 'N/A')}  |  "
        f"Location: {state.get('location', 'N/A')}  |  "
        f"Severity: {state.get('severity', 'N/A')}"
    )
    meta.runs[0].font.size = Pt(9)
    meta.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Parse and write report content
    for line in report_content.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
        # Section headings like "1. Problem Description"
        if re.match(r"^\d+\.", line) and len(line) < 80:
            doc.add_heading(line, level=1)
        # Field lines like "Equipment type: centrifugal pump"
        elif ":" in line:
            field, _, value = line.partition(":")
            p = doc.add_paragraph()
            run_field = p.add_run(field.strip() + ": ")
            run_field.bold = True
            p.add_run(value.strip())
        else:
            doc.add_paragraph(line)

    doc.save(str(report_path))
    return report_path